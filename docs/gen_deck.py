#!/usr/bin/env python3
"""Generate TraceMind business pitch deck."""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
import os

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

# Color palette — dark modern theme
BG_DARK = RGBColor(0x0F, 0x17, 0x2A)
BG_CARD = RGBColor(0x1E, 0x29, 0x3B)
ACCENT = RGBColor(0x38, 0xBD, 0xF8)   # cyan
ACCENT2 = RGBColor(0x22, 0xD3, 0xEE)
GREEN = RGBColor(0x4A, 0xDE, 0x80)
ORANGE = RGBColor(0xFB, 0xBF, 0x24)
RED = RGBColor(0xF8, 0x71, 0x71)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
GRAY = RGBColor(0x94, 0xA3, 0xB8)
LIGHT = RGBColor(0xE2, 0xE8, 0xF0)


def dark_bg(slide):
    bg = slide.background
    fill = bg.fill
    fill.solid()
    fill.fore_color.rgb = BG_DARK


def add_text(slide, text, left, top, width, height, font_size=18, color=WHITE,
             bold=False, align=PP_ALIGN.LEFT, font_name="Helvetica"):
    txBox = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(font_size)
    p.font.color.rgb = color
    p.font.bold = bold
    p.font.name = font_name
    p.alignment = align
    return tf


def add_bullets(slide, items, left, top, width, height, font_size=16, color=LIGHT, spacing=Pt(8)):
    txBox = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = txBox.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        if i == 0:
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()
        p.text = item
        p.font.size = Pt(font_size)
        p.font.color.rgb = color
        p.font.name = "Helvetica"
        p.space_after = spacing
        p.level = 0
        pPr = p._pPr
        if pPr is None:
            from pptx.oxml.ns import qn
            pPr = p._p.get_or_add_pPr()
        from pptx.oxml.ns import qn
        buChar = pPr.makeelement(qn('a:buChar'), {'char': '\u2022'})
        # remove existing bullet
        for existing in pPr.findall(qn('a:buChar')):
            pPr.remove(existing)
        for existing in pPr.findall(qn('a:buNone')):
            pPr.remove(existing)
        pPr.append(buChar)
    return tf


def add_card(slide, left, top, width, height, fill_color=BG_CARD):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(top),
        Inches(width), Inches(height)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill_color
    shape.line.fill.background()
    return shape


def accent_line(slide, left, top, width=2.0):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(left), Inches(top),
        Inches(width), Pt(4)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = ACCENT
    shape.line.fill.background()


# ========== SLIDE 1: TITLE ==========
s = prs.slides.add_slide(prs.slide_layouts[6])  # blank
dark_bg(s)
accent_line(s, 1.0, 2.5, 4.0)
add_text(s, "TraceMind", 1.0, 2.7, 11, 1.2, font_size=54, bold=True, color=WHITE)
add_text(s, "The Memory OS for AI", 1.0, 3.7, 11, 0.8, font_size=32, color=ACCENT)
add_text(s, "Your AI finally remembers. Locally. Privately. Auditably.", 1.0, 4.7, 11, 0.6, font_size=20, color=GRAY)
add_text(s, "Aaditya Srivathsan  |  March 2026", 1.0, 6.2, 11, 0.4, font_size=14, color=GRAY)

# ========== SLIDE 2: THE PROBLEM ==========
s = prs.slides.add_slide(prs.slide_layouts[6])
dark_bg(s)
add_text(s, "The Problem", 0.8, 0.4, 11, 0.8, font_size=36, bold=True, color=WHITE)
accent_line(s, 0.8, 1.1, 2.0)

problems = [
    ("Stateless AI", "Every session starts from zero. Claude Code rediscovers your codebase every single time.", RED),
    ("Endless Repetition", "Users repeat preferences, context, and decisions to AI tools hundreds of times.", ORANGE),
    ("No Audit Trail", "In regulated environments, there is zero accountability for AI-assisted decisions.", RED),
    ("Cloud Lock-in", "Existing memory solutions (mem0, Zep) require sending private data to third-party servers.", ORANGE),
    ("Blind Recording", "Rewind/Limitless record screens but build no structured understanding.", GRAY),
]

for i, (title, desc, col) in enumerate(problems):
    y = 1.6 + i * 1.1
    add_card(s, 0.8, y, 11.7, 0.95)
    add_text(s, title, 1.1, y + 0.08, 2.5, 0.4, font_size=18, bold=True, color=col)
    add_text(s, desc, 3.6, y + 0.08, 8.5, 0.8, font_size=15, color=LIGHT)

# ========== SLIDE 3: THE SOLUTION ==========
s = prs.slides.add_slide(prs.slide_layouts[6])
dark_bg(s)
add_text(s, "TraceMind", 0.8, 0.4, 11, 0.8, font_size=36, bold=True, color=ACCENT)
add_text(s, "A local-only memory OS that makes AI remember, learn, and explain.", 0.8, 1.1, 11, 0.5, font_size=20, color=GRAY)
accent_line(s, 0.8, 1.7, 3.0)

features = [
    ("Capture", "Watches your browser, Claude Code, clipboard\u2014with explicit consent"),
    ("Structure", "Builds entity graphs and semantic clusters,\nnot raw recordings"),
    ("Learn", "RL-based policy that improves retrieval\nover time from your feedback"),
    ("Audit", "Full Merkle-chain trace for every\nAI decision. Deterministic replay."),
    ("Local", "100% on-device. No cloud. No telemetry.\nYour data never leaves your machine."),
]

for i, (title, desc) in enumerate(features):
    x = 0.8 + i * 2.4
    add_card(s, x, 2.2, 2.2, 4.5)
    add_text(s, title, x + 0.15, 2.4, 1.9, 0.5, font_size=22, bold=True, color=ACCENT)
    add_text(s, desc, x + 0.15, 3.1, 1.9, 3.2, font_size=14, color=LIGHT)

# ========== SLIDE 4: HOW IT WORKS ==========
s = prs.slides.add_slide(prs.slide_layouts[6])
dark_bg(s)
add_text(s, "How It Works", 0.8, 0.4, 11, 0.8, font_size=36, bold=True, color=WHITE)
accent_line(s, 0.8, 1.1, 2.0)

pipeline = [
    ("Capture", "Browser, clipboard,\nClaude Code hooks", ACCENT),
    ("\u2192", "", GRAY),
    ("Canonicalize", "Entity extraction,\ntriple generation", ACCENT2),
    ("\u2192", "", GRAY),
    ("Govern", "PII filter, schema\nvalidation, user rules", GREEN),
    ("\u2192", "", GRAY),
    ("Store", "Graph + vectors +\nclusters + traces", ACCENT),
    ("\u2192", "", GRAY),
    ("Retrieve", "3-phase recall\nwith bounded budget", ACCENT2),
    ("\u2192", "", GRAY),
    ("Learn", "UCB bandit \u2192 RL\npolicy from trajectories", GREEN),
]

x = 0.3
for title, desc, col in pipeline:
    if title == "\u2192":
        add_text(s, "\u2192", x, 2.8, 0.5, 0.5, font_size=28, color=GRAY, align=PP_ALIGN.CENTER)
        x += 0.5
    else:
        add_card(s, x, 2.2, 1.7, 2.5)
        add_text(s, title, x + 0.1, 2.35, 1.5, 0.4, font_size=18, bold=True, color=col)
        add_text(s, desc, x + 0.1, 2.9, 1.5, 1.5, font_size=13, color=LIGHT)
        x += 1.9

# Bottom bar: constraints
add_card(s, 0.8, 5.3, 11.7, 1.2)
add_text(s, "All on-device  \u2022  Rust core  \u2022  <200MB RAM  \u2022  <250MB install  \u2022  macOS / Windows / Linux",
         1.1, 5.5, 11, 0.8, font_size=18, color=ACCENT, align=PP_ALIGN.CENTER)

# ========== SLIDE 5: CLAUDE CODE INTEGRATION ==========
s = prs.slides.add_slide(prs.slide_layouts[6])
dark_bg(s)
add_text(s, "Claude Code + TraceMind", 0.8, 0.4, 11, 0.8, font_size=36, bold=True, color=WHITE)
accent_line(s, 0.8, 1.1, 2.5)

# Left: Today
add_card(s, 0.8, 1.6, 5.5, 5.0)
add_text(s, "Today: Stateless", 1.1, 1.8, 5.0, 0.5, font_size=22, bold=True, color=RED)
add_bullets(s, [
    "Forgets everything between sessions",
    "Rediscovers your codebase every time",
    "Can't learn from past mistakes",
    "No audit trail of AI code decisions",
], 1.1, 2.5, 4.8, 3.5, font_size=16, color=GRAY)

# Right: With TraceMind
add_card(s, 7.0, 1.6, 5.5, 5.0)
add_text(s, "With TraceMind", 7.3, 1.8, 5.0, 0.5, font_size=22, bold=True, color=GREEN)
add_bullets(s, [
    "Session 1: generic",
    "Session 10: knows your naming conventions",
    "Session 50: knows what succeeded / failed",
    "Session 200: deep project knowledge graph",
    "MCP server + passive hooks = bidirectional",
], 7.3, 2.5, 4.8, 3.8, font_size=16, color=LIGHT)

# ========== SLIDE 6: WHY NOW ==========
s = prs.slides.add_slide(prs.slide_layouts[6])
dark_bg(s)
add_text(s, "Why Now", 0.8, 0.4, 11, 0.8, font_size=36, bold=True, color=WHITE)
accent_line(s, 0.8, 1.1, 2.0)

reasons = [
    ("86%", "of enterprises will deploy\nAI agents by 2027", "Gartner"),
    ("\u20AC35M", "max fines under EU AI Act\nfor non-compliance", "EU AI Act 2024"),
    ("$21.1B", "Explainable AI market\nby 2030 (from $7.8B)", "Grand View Research"),
    ("$5.8B", "AI Governance market\nby 2029 (45% CAGR)", "MarketsAndMarkets"),
]

for i, (stat, desc, source) in enumerate(reasons):
    x = 0.8 + i * 3.1
    add_card(s, x, 1.6, 2.9, 3.5)
    add_text(s, stat, x + 0.2, 1.9, 2.5, 0.8, font_size=36, bold=True, color=ACCENT)
    add_text(s, desc, x + 0.2, 2.8, 2.5, 1.2, font_size=15, color=LIGHT)
    add_text(s, source, x + 0.2, 4.2, 2.5, 0.4, font_size=11, color=GRAY)

add_card(s, 0.8, 5.5, 11.7, 1.2)
add_text(s, "Microsoft Recall backlash proved: users WANT memory but DEMAND privacy.\nNo product combines: local + structured + learning + auditable. Until now.",
         1.1, 5.6, 11.2, 1.0, font_size=17, color=ORANGE, align=PP_ALIGN.CENTER)

# ========== SLIDE 7: WOW MOMENTS ==========
s = prs.slides.add_slide(prs.slide_layouts[6])
dark_bg(s)
add_text(s, "Product Moments That Matter", 0.8, 0.4, 11, 0.8, font_size=36, bold=True, color=WHITE)
accent_line(s, 0.8, 1.1, 3.0)

moments = [
    ("\u201cIt remembered what I browsed 2 weeks ago\nand surfaced it in Claude Code.\u201d", ACCENT),
    ("\u201cShow me WHY my AI made that decision.\u201d\n\u2192 Full trace replay with every memory ID.", GREEN),
    ("\u201cAfter 50 sessions, retrieval quality\nmeasurably improved.\u201d", ACCENT2),
    ("\u201cAll my data is on MY machine.\nI can see exactly what it knows.\u201d", ORANGE),
]

for i, (text, col) in enumerate(moments):
    y = 1.6 + i * 1.4
    add_card(s, 0.8, y, 11.7, 1.2)
    add_text(s, f"{i+1}", 1.1, y + 0.15, 0.5, 0.8, font_size=32, bold=True, color=col)
    add_text(s, text, 1.8, y + 0.1, 10.2, 1.0, font_size=18, color=LIGHT)

# ========== SLIDE 8: MARKET OPPORTUNITY ==========
s = prs.slides.add_slide(prs.slide_layouts[6])
dark_bg(s)
add_text(s, "Market Opportunity", 0.8, 0.4, 11, 0.8, font_size=36, bold=True, color=WHITE)
accent_line(s, 0.8, 1.1, 2.0)

markets = [
    ("TAM", "$80B+", "AI Governance ($5.8B) + Legal AI ($3.9B)\n+ Compliance Software ($75B by 2032)", ACCENT),
    ("SAM", "$8B", "Developer tools + regulated enterprise\nAI users in North America & Europe", ACCENT2),
    ("SOM (Y1)", "$7M", "50K consumers @ $10/mo\n+ 20 enterprise @ $50K/yr", GREEN),
]

for i, (label, amount, desc, col) in enumerate(markets):
    x = 0.8 + i * 4.1
    add_card(s, x, 1.6, 3.8, 3.0)
    add_text(s, label, x + 0.2, 1.8, 3.4, 0.4, font_size=16, bold=True, color=GRAY)
    add_text(s, amount, x + 0.2, 2.2, 3.4, 0.7, font_size=40, bold=True, color=col)
    add_text(s, desc, x + 0.2, 3.0, 3.4, 1.2, font_size=14, color=LIGHT)

add_text(s, "Wedge:  Consumer desktop app  \u2192  Claude Code power users  \u2192  Enterprise governance",
         0.8, 5.2, 11.7, 0.5, font_size=18, color=ORANGE, align=PP_ALIGN.CENTER)

# ========== SLIDE 9: COMPETITIVE LANDSCAPE ==========
s = prs.slides.add_slide(prs.slide_layouts[6])
dark_bg(s)
add_text(s, "Competitive Landscape", 0.8, 0.4, 11, 0.8, font_size=36, bold=True, color=WHITE)
accent_line(s, 0.8, 1.1, 2.5)

# Table header
headers = ["", "Local-Only", "Structured\nMemory", "Audit\nTrail", "Self-\nImproving", "Lightweight", "Open"]
rows = [
    ("TraceMind", [True, True, True, True, True, True]),
    ("Rewind / Limitless", [True, False, False, False, False, False]),
    ("mem0 / Zep", [False, False, False, False, True, True]),
    ("Apple Intelligence", [True, False, False, False, True, False]),
    ("Microsoft Recall", [True, False, False, False, False, False]),
]

# Header row
for j, h in enumerate(headers):
    x = 0.8 + j * 1.7
    add_card(s, x, 1.5, 1.6, 0.8)
    add_text(s, h, x + 0.05, 1.55, 1.5, 0.7, font_size=12, bold=True, color=ACCENT, align=PP_ALIGN.CENTER)

# Data rows
for i, (name, vals) in enumerate(rows):
    y = 2.4 + i * 0.9
    # Name
    col = GREEN if i == 0 else LIGHT
    add_card(s, 0.8, y, 1.6, 0.75, fill_color=RGBColor(0x16, 0x1F, 0x33) if i > 0 else RGBColor(0x0D, 0x3B, 0x2E))
    add_text(s, name, 0.85, y + 0.1, 1.5, 0.5, font_size=12, bold=(i == 0), color=col, align=PP_ALIGN.CENTER)
    for j, v in enumerate(vals):
        x = 2.5 + j * 1.7
        bg = RGBColor(0x16, 0x1F, 0x33) if not v else (RGBColor(0x0D, 0x3B, 0x2E) if i == 0 else RGBColor(0x1E, 0x29, 0x3B))
        add_card(s, x, y, 1.6, 0.75, fill_color=bg)
        symbol = "\u2713" if v else "\u2717"
        sym_color = GREEN if v else RGBColor(0x4B, 0x55, 0x63)
        add_text(s, symbol, x + 0.05, y + 0.08, 1.5, 0.5, font_size=20, color=sym_color, align=PP_ALIGN.CENTER)

# ========== SLIDE 10: TECHNICAL MOAT ==========
s = prs.slides.add_slide(prs.slide_layouts[6])
dark_bg(s)
add_text(s, "Technical Moat", 0.8, 0.4, 11, 0.8, font_size=36, bold=True, color=WHITE)
accent_line(s, 0.8, 1.1, 2.0)

moats = [
    ("Rust Data Plane", "5x less RAM than Python.\nNo GC pauses. <50MB binary.", ACCENT),
    ("Embedded DBs", "Kuzu (graph) + LanceDB (vector).\nNo server processes. <200MB total.", ACCENT2),
    ("Merkle Audit Log", "Cryptographic hash chain.\nTamper-evident. Forensic-grade.", GREEN),
    ("RL Learning", "Memory-R1 + Graph-R1 trajectories.\nBandit \u2192 policy net \u2192 GRPO.", ORANGE),
    ("FSM Execution", "Deterministic state machine.\nBitwise reproducible decisions.", ACCENT),
    ("Cross-Platform", "Tauri: one codebase.\nmacOS, Windows, Linux.", ACCENT2),
]

for i, (title, desc, col) in enumerate(moats):
    row, c = divmod(i, 3)
    x = 0.8 + c * 4.1
    y = 1.6 + row * 2.6
    add_card(s, x, y, 3.8, 2.2)
    add_text(s, title, x + 0.2, y + 0.15, 3.4, 0.5, font_size=20, bold=True, color=col)
    add_text(s, desc, x + 0.2, y + 0.75, 3.4, 1.2, font_size=15, color=LIGHT)

# ========== SLIDE 11: BUSINESS MODEL ==========
s = prs.slides.add_slide(prs.slide_layouts[6])
dark_bg(s)
add_text(s, "Business Model", 0.8, 0.4, 11, 0.8, font_size=36, bold=True, color=WHITE)
accent_line(s, 0.8, 1.1, 2.0)

# Tiers
tiers = [
    ("Free", "$0", "Local memory (30-day window)\nBasic graph visualization\nCLI access", GRAY),
    ("Pro", "$10/mo", "Unlimited memory retention\nRL learning + trajectory training\nClaude Code MCP integration\nChrome extension\nFull audit trail", ACCENT),
    ("Enterprise", "$50\u2013200K/yr", "Docker deployment\nTeam knowledge graphs\nSOC2 compliance toolkit\nPolicy-as-Code engine\nMerkle audit + export\nSSO / RBAC", GREEN),
]

for i, (name, price, features, col) in enumerate(tiers):
    x = 0.8 + i * 4.1
    add_card(s, x, 1.5, 3.8, 5.2)
    add_text(s, name, x + 0.2, 1.7, 3.4, 0.5, font_size=24, bold=True, color=col)
    add_text(s, price, x + 0.2, 2.2, 3.4, 0.5, font_size=28, bold=True, color=WHITE)
    add_text(s, features, x + 0.2, 3.0, 3.4, 3.5, font_size=14, color=LIGHT)

# ========== SLIDE 12: GO TO MARKET ==========
s = prs.slides.add_slide(prs.slide_layouts[6])
dark_bg(s)
add_text(s, "Go-to-Market", 0.8, 0.4, 11, 0.8, font_size=36, bold=True, color=WHITE)
accent_line(s, 0.8, 1.1, 2.0)

phases = [
    ("Q2 2026", "Developer Launch", "Desktop app + Claude Code MCP\nTarget: developers, power users\nHomebrew, Product Hunt", ACCENT),
    ("Q3 2026", "Consumer Growth", "Chrome extension + marketing\nTarget: 50K users\nContent marketing, community", ACCENT2),
    ("Q4 2026", "Enterprise", "Docker + compliance toolkit\nTarget: legal, finance, healthcare\n20 enterprise contracts", GREEN),
    ("2027", "Platform", "API marketplace\nTeam knowledge graphs\nSeries A", ORANGE),
]

for i, (time, title, desc, col) in enumerate(phases):
    x = 0.5 + i * 3.2
    add_card(s, x, 1.6, 3.0, 4.5)
    add_text(s, time, x + 0.15, 1.75, 2.7, 0.4, font_size=14, bold=True, color=col)
    add_text(s, title, x + 0.15, 2.15, 2.7, 0.5, font_size=20, bold=True, color=WHITE)
    add_text(s, desc, x + 0.15, 2.8, 2.7, 2.8, font_size=14, color=LIGHT)

# Distribution
add_card(s, 0.8, 6.4, 11.7, 0.8)
add_text(s, "Distribution:  Claude Code Marketplace  \u2022  Homebrew  \u2022  Product Hunt  \u2022  Dev Communities  \u2022  Direct Sales",
         1.1, 6.45, 11.2, 0.6, font_size=16, color=ACCENT, align=PP_ALIGN.CENTER)

# ========== SLIDE 13: TEAM ==========
s = prs.slides.add_slide(prs.slide_layouts[6])
dark_bg(s)
add_text(s, "Team", 0.8, 0.4, 11, 0.8, font_size=36, bold=True, color=WHITE)
accent_line(s, 0.8, 1.1, 2.0)

add_card(s, 2.5, 2.0, 8.3, 3.5)
add_text(s, "Aaditya Srivathsan", 2.8, 2.3, 7.7, 0.6, font_size=28, bold=True, color=WHITE)
add_text(s, "Founder & CEO", 2.8, 2.9, 7.7, 0.4, font_size=18, color=ACCENT)
add_text(s, "Software engineer with deep expertise in systems programming (Rust, Python),\n"
         "ML/AI (reinforcement learning, JEPA, agent architectures), and distributed systems.\n\n"
         "Previously built AgentMem v2 \u2014 a self-improving agent memory system\n"
         "with knowledge graphs, bandit-based retrieval, and multi-agent reasoning.\n\n"
         "Passionate about privacy-first AI infrastructure.",
         2.8, 3.5, 7.7, 2.5, font_size=15, color=LIGHT)

# ========== SLIDE 14: THE ASK ==========
s = prs.slides.add_slide(prs.slide_layouts[6])
dark_bg(s)
add_text(s, "The Ask", 0.8, 0.4, 11, 0.8, font_size=36, bold=True, color=WHITE)
accent_line(s, 0.8, 1.1, 2.0)

add_card(s, 0.8, 1.6, 5.5, 2.5)
add_text(s, "Seed Round", 1.1, 1.8, 5.0, 0.5, font_size=22, bold=True, color=ACCENT)
add_text(s, "[Amount TBD]", 1.1, 2.3, 5.0, 0.5, font_size=28, bold=True, color=WHITE)

add_card(s, 7.0, 1.6, 5.5, 2.5)
add_text(s, "Use of Funds", 7.3, 1.8, 5.0, 0.5, font_size=22, bold=True, color=ACCENT)
add_text(s, "Engineering: 60%\nGo-to-Market: 25%\nOperations: 15%", 7.3, 2.4, 5.0, 1.5, font_size=18, color=LIGHT)

add_card(s, 0.8, 4.5, 11.7, 2.2)
add_text(s, "Key Milestones", 1.1, 4.7, 11.2, 0.5, font_size=22, bold=True, color=GREEN)
add_bullets(s, [
    "50,000 consumer users within 6 months of launch",
    "20 enterprise contracts within 12 months",
    "Demonstrated RL-based retrieval improvement (published benchmarks)",
    "Series A readiness by Q1 2027",
], 1.1, 5.2, 11.0, 1.5, font_size=16, color=LIGHT)

# ========== SLIDE 15: CLOSING ==========
s = prs.slides.add_slide(prs.slide_layouts[6])
dark_bg(s)
accent_line(s, 1.0, 2.5, 4.0)
add_text(s, "TraceMind", 1.0, 2.7, 11, 1.2, font_size=54, bold=True, color=WHITE)
add_text(s, "Your AI finally remembers.", 1.0, 3.8, 11, 0.6, font_size=28, color=ACCENT)
add_text(s, "Local. Private. Auditable. Self-improving.", 1.0, 4.5, 11, 0.5, font_size=20, color=GRAY)
add_text(s, "aaditya@tracemind.ai", 1.0, 5.8, 11, 0.4, font_size=18, color=LIGHT)

# Save
outpath = os.path.join(os.path.dirname(__file__), "03-business-deck.pptx")
prs.save(outpath)
print(f"Saved to {outpath}")
